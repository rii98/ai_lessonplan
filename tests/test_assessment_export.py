"""Every question type, rendered to every format from the one shared layout."""

from __future__ import annotations

import io

import pytest
from docx import Document
from pptx import Presentation

from lessonforge.domain.artifacts import Quiz, Worksheet
from lessonforge.domain.assessment import AssessmentSpec, TypeSpec, view
from lessonforge.domain.ldd import Question
from lessonforge.export import ArtifactKind, build_renderer, formats_for

_REF = {"board": "CDC", "grade": 7, "subject": "Science", "code": None}
_OBJ = [{"id": "O1", "statement": "Explain that sound is a vibration", "bloom": "understand"}]


def _questions():
    base = {"objective_ids": ["O1"]}
    return [Question.model_validate({**base, **q}) for q in [
        {"id": "Q1", "type": "mcq", "prompt": "Which makes sound?", "options": ["stone", "string"],
         "answer": "string", "difficulty": "easy"},
        {"id": "Q2", "type": "true_false", "prompt": "Sound travels in space.", "answer": "False"},
        {"id": "Q3", "type": "fill_blank", "prompt": "Sound is a _____.", "answer": "vibration"},
        {"id": "Q4", "type": "matching", "prompt": "Match", "pairs": [
            {"left": "Madal", "right": "skin"}, {"left": "Flute", "right": "air"},
            {"left": "Sarangi", "right": "string"}]},
        {"id": "Q5", "type": "short_answer", "prompt": "Name a drum.", "answer": "madal",
         "key_points": ["names one"]},
        {"id": "Q6", "type": "long_answer", "prompt": "Explain why space is silent.",
         "answer": "No medium to carry vibrations.", "difficulty": "hard",
         "key_points": ["needs a medium", "space is a vacuum"]},
        {"id": "Q7", "type": "numerical", "prompt": "120 km in 4 h: speed?", "answer": "30 km/h",
         "key_points": ["120 ÷ 4"]},
        {"id": "Q8", "type": "ordering", "prompt": "Arrange the steps.",
         "options": ["vibrate", "air moves", "ear hears"]},
    ]]


def _spec():
    return AssessmentSpec(types=[
        TypeSpec(type="mcq", count=1, marks_each=1), TypeSpec(type="true_false", count=1, marks_each=1),
        TypeSpec(type="fill_blank", count=1, marks_each=1), TypeSpec(type="matching", count=1, marks_each=3),
        TypeSpec(type="short_answer", count=1, marks_each=2), TypeSpec(type="long_answer", count=1, marks_each=5),
        TypeSpec(type="numerical", count=1, marks_each=4), TypeSpec(type="ordering", count=1, marks_each=3),
    ])


def _quiz(spec=True):
    return Quiz(topic="Sound", curriculum_ref=_REF, objectives=_OBJ, questions=_questions(),
                spec=_spec() if spec else None)


def _worksheet():
    return Worksheet(topic="Sound", curriculum_ref=_REF, objectives=_OBJ, tasks=["Hum and feel."],
                     questions=_questions(), spec=_spec())


def _docx_text(content: bytes) -> str:
    d = Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def _pptx_text(content: bytes) -> str:
    out = []
    for s in Presentation(io.BytesIO(content)).slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += [c.text for r in sh.table.rows for c in r.cells]
    return "\n".join(out)


def test_quiz_and_worksheet_are_available_as_md_docx_and_pptx():
    for kind in (ArtifactKind.quiz, ArtifactKind.worksheet):
        assert {"md", "docx", "pptx"} <= set(formats_for(kind))


# ── markdown ─────────────────────────────────────────────────────────────────
def test_markdown_sections_marks_and_every_type():
    md = build_renderer(ArtifactKind.quiz, "md").render(_quiz()).content.decode()
    assert "## Section A — Multiple Choice Questions — 1 × 1 = 1 mark" in md
    assert "## Section H — Arrange in the Correct Order" in md
    assert "| Column A | Column B |" in md and "| 1. Madal |" in md
    assert "Sound is a _____." in md
    assert "Correct order:" in md
    assert "**Score:** _____ / 20" in md                      # full marks, not question count
    key = md.split("## Answer Key")[1]
    assert "6. No medium to carry vibrations. _(hard)_" in key
    assert "   - needs a medium" in key
    assert "30 km/h" in key


def test_markdown_key_matches_the_printed_shuffle():
    q = _questions()[3]
    v = view(q, 4)
    md = build_renderer(ArtifactKind.quiz, "md").render(_quiz()).content.decode()
    assert v.answer in md.split("## Answer Key")[1]
    assert v.right[0].split(". ", 1)[1] in md


def test_markdown_is_byte_stable_across_renders():
    r = build_renderer(ArtifactKind.quiz, "md")
    assert r.render(_quiz()).content == r.render(_quiz()).content


def test_without_a_spec_the_flat_legacy_layout_is_kept():
    md = build_renderer(ArtifactKind.quiz, "md").render(_quiz(spec=False)).content.decode()
    assert "## Section" not in md and "**1. Which makes sound?**" in md
    assert "**Score:** _____ / 8" in md


# ── docx ─────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("kind,maker", [(ArtifactKind.quiz, _quiz), (ArtifactKind.worksheet, _worksheet)])
def test_docx_renders_sections_tables_and_key(kind, maker):
    art = build_renderer(kind, "docx").render(maker())
    doc = Document(io.BytesIO(art.content))
    text = _docx_text(art.content)
    assert "Section D — Match the Following" in text and "[3 × 1 = 3 marks]" not in text
    assert len(doc.tables) == 1                                  # the matching table
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Column A", "Column B"]
    assert len(doc.tables[0].rows) == 1 + 3
    assert "Full marks: 20" in text or "/ 20" in text        # worksheet / quiz header
    assert "Answer Key" in text and "needs a medium" in text and "30 km/h" in text
    assert "Correct order:" in text


def test_docx_key_agrees_with_the_printed_matching_column():
    q = _questions()[3]
    v = view(q, 4)
    text = _docx_text(build_renderer(ArtifactKind.quiz, "docx").render(_quiz()).content)
    assert v.answer in text
    table_right = [r.split(". ", 1)[1] for r in v.right]
    assert all(t in text for t in table_right)


# ── pptx ─────────────────────────────────────────────────────────────────────
def test_pptx_quiz_has_one_slide_per_question_dividers_key_and_private_notes():
    art = build_renderer(ArtifactKind.quiz, "pptx").render(_quiz())
    assert art.filename.endswith("_quiz.pptx") and "presentationml" in art.media_type
    prs = Presentation(io.BytesIO(art.content))
    text = _pptx_text(art.content)
    for n in range(1, 9):
        assert f"Question {n}" in text
    assert "Section D — Match the Following" in text and "Column A" in text
    assert "Answer Key" in text and "Sources" in text
    # the answer is in the speaker notes of the question slide, not on its face
    q6 = next(s for s in prs.slides if any(
        sh.has_text_frame and sh.text_frame.text == "Question 6" for sh in s.shapes))
    notes = q6.notes_slide.notes_text_frame.text
    assert "ANSWER: No medium to carry vibrations." in notes and "needs a medium" in notes
    assert "No medium to carry" not in "\n".join(
        sh.text_frame.text for sh in q6.shapes if sh.has_text_frame)


def test_pptx_worksheet_leads_with_objectives_and_tasks():
    text = _pptx_text(build_renderer(ArtifactKind.worksheet, "pptx").render(_worksheet()).content)
    assert text.index("Objectives") < text.index("Tasks") < text.index("Question 1")


def test_pptx_answer_key_paginates_long_keys():
    long = [Question(id=f"Q{i}", type="long_answer", prompt="p", answer="word " * 80,
                     objective_ids=["O1"]) for i in range(1, 9)]
    quiz = Quiz(topic="T", curriculum_ref=_REF, objectives=_OBJ, questions=long)
    text = _pptx_text(build_renderer(ArtifactKind.quiz, "pptx").render(quiz).content)
    assert "Answer Key (1/" in text


def test_pptx_renders_a_legacy_flat_quiz():
    art = build_renderer(ArtifactKind.quiz, "pptx").render(_quiz(spec=False))
    assert "Question 8" in _pptx_text(art.content)
