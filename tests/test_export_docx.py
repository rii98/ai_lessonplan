"""DOCX renderers — verified structurally (open with python-docx, check the 5E
table, headings, answer key) and for the Devanagari de-risk: codepoints survive
and the complex-script font hint is set on runs."""

from __future__ import annotations

import io
import zipfile

from docx import Document

from lessonforge.export import ArtifactKind, ExportOptions, build_renderer


def _open(content: bytes) -> Document:
    return Document(io.BytesIO(content))


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_lesson_plan_opens_and_has_5e_table(export_ldd):
    art = build_renderer(ArtifactKind.lesson_plan, "docx").render(export_ldd)
    assert art.media_type.endswith("wordprocessingml.document")
    assert art.filename.endswith(".docx")
    doc = _open(art.content)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    # header row + one row per phase
    assert len(table.rows) == 1 + len(export_ldd.phases)
    headers = [c.text for c in table.rows[0].cells]
    assert headers == ["Phase", "Teacher Activities", "Student Activities", "Time"]
    text = _all_text(doc)
    assert "Learning Objectives" in text
    assert "Homework" in text
    assert "planned 45 min" in text  # timing surfaced


def test_devanagari_codepoints_and_cs_font_survive(export_ldd):
    art = build_renderer(ArtifactKind.lesson_plan, "docx").render(export_ldd)
    xml = zipfile.ZipFile(io.BytesIO(art.content)).read("word/document.xml").decode("utf-8")
    assert "संलग्न" in xml                     # Devanagari intact
    assert 'w:cs="Noto Sans Devanagari"' in xml  # complex-script font hint set


def test_devanagari_font_is_configurable(export_ldd):
    opts = ExportOptions(devanagari_font="Mangal")
    art = build_renderer(ArtifactKind.lesson_plan, "docx", opts).render(export_ldd)
    xml = zipfile.ZipFile(io.BytesIO(art.content)).read("word/document.xml").decode("utf-8")
    assert 'w:cs="Mangal"' in xml


def test_worksheet_has_questions_and_answer_key_on_new_page(export_ldd):
    art = build_renderer(ArtifactKind.worksheet, "docx").render(export_ldd)
    doc = _open(art.content)
    text = _all_text(doc)
    assert "Practice Questions" in text
    assert "Answer Key" in text
    assert "goat" in text  # mcq option + answer
    # a page break separates the questions from the key
    xml = zipfile.ZipFile(io.BytesIO(art.content)).read("word/document.xml").decode("utf-8")
    assert 'w:type="page"' in xml


def test_quiz_opens_and_carries_answer_key(export_ldd):
    art = build_renderer(ArtifactKind.quiz, "docx").render(export_ldd)
    text = _all_text(_open(art.content))
    assert "Quiz — Components of Environment" in text
    assert "Answer Key" in text
    assert "rice plant" in text  # short-answer answer in the key
