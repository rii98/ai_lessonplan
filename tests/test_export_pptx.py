"""PPTX renderer — hook-first ordering, one slide per phase, Devanagari shaping."""

from __future__ import annotations

import io
import zipfile

from pptx import Presentation

from lessonforge.export import ArtifactKind, build_renderer


def _slides_text(content: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(content))
    out = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        out.append("\n".join(texts))
    return out


def test_deck_opens_and_is_hook_first(export_ldd):
    art = build_renderer(ArtifactKind.slides, "pptx").render(export_ldd)
    assert art.media_type.endswith("presentationml.presentation")
    slides = _slides_text(art.content)
    # title, hook, objectives, 3 phases, check, + mandatory Sources slide
    assert len(slides) == 3 + len(export_ldd.phases) + 1 + 1
    assert export_ldd.topic in slides[0]              # title first
    assert export_ldd.engagement_hook.prompt in slides[1]   # hook BEFORE objectives
    assert "What we'll be able to do" in slides[2]    # objectives after the hook
    assert "Sources" in slides[-1]                    # citations last, mandatory


def test_every_phase_becomes_a_slide(export_ldd):
    slides = _slides_text(build_renderer(ArtifactKind.slides, "pptx").render(export_ldd).content)
    joined = "\n".join(slides)
    for phase in export_ldd.phases:
        assert phase.name_en in joined


def test_devanagari_shaped_with_cs_font(export_ldd):
    art = build_renderer(ArtifactKind.slides, "pptx").render(export_ldd)
    zf = zipfile.ZipFile(io.BytesIO(art.content))
    xml = "".join(zf.read(n).decode("utf-8") for n in zf.namelist() if n.endswith(".xml"))
    assert "संलग्न" in xml
    assert 'typeface="Noto Sans Devanagari"' in xml
