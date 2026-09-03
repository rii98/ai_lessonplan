"""Inline rich-text for the Office exporters: LaTeX math is transliterated to
Unicode and Markdown emphasis becomes real runs, so ``$``/``**`` never leak into
a .docx / .pptx as literal characters (the same problem the chat renderer solves
in the browser, one layer down)."""

from __future__ import annotations

import io

from lessonforge.export.richtext import Segment, latex_to_unicode, parse_inline

# ── LaTeX → Unicode ──────────────────────────────────────────────────────────

def test_superscripts_and_subscripts():
    assert latex_to_unicode("x^2 = 16") == "x² = 16"
    assert latex_to_unicode("a_1 + b_2 = c_3") == "a₁ + b₂ = c₃"


def test_fraction_and_sqrt():
    out = latex_to_unicode(r"\frac{-b \pm \sqrt{b^2-4ac}}{2a}")
    assert out == "(-b ± √(b²-4ac))/2a"


def test_symbols_and_greek():
    assert latex_to_unicode(r"x \le 5 \ge 2 \times 3 \div 4") == "x ≤ 5 ≥ 2 × 3 ÷ 4"
    assert latex_to_unicode(r"\alpha + \pi r^2") == "α + π r²"


def test_unknown_command_degrades_to_readable_source():
    # no raw backslashes survive; the command name is kept as legible text
    out = latex_to_unicode(r"\foo x")
    assert "\\" not in out and "foo" in out


# ── inline Markdown split ────────────────────────────────────────────────────

def test_emphasis_becomes_segments():
    segs = parse_inline("plain **bold** and *italic* and `code`")
    assert Segment("bold", bold=True) in segs
    assert Segment("italic", italic=True) in segs
    assert Segment("code", code=True) in segs


def test_math_is_extracted_and_marked_italic():
    segs = parse_inline("Solve $x^2 = 16$ now")
    math = [s for s in segs if s.text == "x² = 16"]
    assert math and math[0].italic and not math[0].code


def test_currency_is_not_treated_as_math():
    segs = parse_inline("It costs $5 and $10 total.")
    joined = "".join(s.text for s in segs)
    assert joined == "It costs $5 and $10 total."


def test_display_and_backslash_delimiters():
    for src in (r"$$E = mc^2$$", r"\(E = mc^2\)", r"\[E = mc^2\]"):
        segs = parse_inline(src)
        assert any("E = mc²" in s.text for s in segs), src


def test_empty_and_plain():
    assert parse_inline("") == []
    assert parse_inline("just text") == [Segment("just text")]


# ── end-to-end through the renderers ─────────────────────────────────────────

def _docx_runs(export_ldd):
    from docx import Document

    from lessonforge.export.base import ExportOptions
    from lessonforge.export.docx_render import _run

    doc = Document()
    p = doc.add_paragraph()
    _run(p, r"Solve $x^2=16$ using **factorization** and `solve()`.", ExportOptions())
    return p.runs


def test_docx_run_splits_styles_and_renders_math(export_ldd):
    runs = _docx_runs(export_ldd)
    text = "".join(r.text for r in runs)
    assert "x²=16" in text
    assert "$" not in text and "**" not in text and "`" not in text
    assert any(r.bold and "factorization" in r.text for r in runs)
    assert any(r.font.name == "Consolas" and r.text == "solve()" for r in runs)


def test_pptx_bullets_render_math_and_emphasis():
    from pptx import Presentation

    from lessonforge.domain.artifacts import Slide, Slides
    from lessonforge.domain.ldd import CurriculumRef
    from lessonforge.export import ArtifactKind, build_renderer

    deck = Slides(
        topic="Quadratics",
        curriculum_ref=CurriculumRef(board="NEB", grade=10, subject="Math"),
        subtitle="Grade 10",
        grounding_sources=["math_10.md"],
        slides=[Slide(heading="Formula", bullets=[
            r"Roots: $x = \frac{-b \pm \sqrt{b^2-4ac}}{2a}$",
            r"Need $b^2 \ge 0$ and **factorise**",
        ])],
    )
    art = build_renderer(ArtifactKind.slides, "pptx").render(deck)
    prs = Presentation(io.BytesIO(art.content))
    text = "\n".join(
        "".join(r.text for r in para.runs)
        for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
        for para in shape.text_frame.paragraphs
    )
    assert "(-b ± √(b²-4ac))/2a" in text
    assert "b² ≥ 0" in text
    assert "$" not in text and "**" not in text
